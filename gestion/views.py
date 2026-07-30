from django.db.models import Q
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from django.shortcuts import render, redirect, get_object_or_404
from django.http import HttpResponse, JsonResponse
from django.template.loader import render_to_string
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.contrib.auth import logout
from django.db.models import Sum, Count
from .models import Proveedor, EmpresaHolding, RequerimientoTicket, RegistroCompraPago, PagoProveedorAdjudicado, Inventario

def custom_logout(request):
    logout(request)
    return redirect('login')

@login_required
def dashboard(request):
    total_proveedores = Proveedor.objects.count()
    total_holding = EmpresaHolding.objects.count()
    total_compras = RegistroCompraPago.objects.count()
    total_inventario = Inventario.objects.count()
    
    req_pendientes = RequerimientoTicket.objects.filter(estatus__in=['PENDIENTE', 'EN_CURSO']).order_by('-id')
    ultimas_compras = RegistroCompraPago.objects.prefetch_related('pagos_proveedores__proveedor').select_related('empresa_destino').order_by('-id')[:10]

    # Datos para Gráficos Dinámicos
    # 1. Gastos Totales por Empresa Holding
    gastos_empresa = (
        PagoProveedorAdjudicado.objects
        .values('solicitud__empresa_destino__nombre')
        .annotate(total=Sum('monto_pagado'))
        .order_by('-total')
    )
    labels_gastos = [g['solicitud__empresa_destino__nombre'] or "Sin Asignar" for g in gastos_empresa]
    data_gastos = [float(g['total'] or 0.0) for g in gastos_empresa]

    # 2. Resumen de Inventario por Estatus
    inv_estatus = (
        Inventario.objects
        .values('estatus')
        .annotate(total=Count('id'))
    )
    dict_estatus = dict(Inventario.ESTATUS_INVENTARIO)
    labels_inv = [dict_estatus.get(i['estatus'], i['estatus']) for i in inv_estatus]
    data_inv = [i['total'] for i in inv_estatus]

    context = {
        'total_proveedores': total_proveedores,
        'total_holding': total_holding,
        'total_compras': total_compras,
        'total_inventario': total_inventario,
        'req_pendientes': req_pendientes,
        'ultimas_compras': ultimas_compras,
        'labels_gastos': labels_gastos,
        'data_gastos': data_gastos,
        'labels_inv': labels_inv,
        'data_inv': data_inv,
    }
    return render(request, 'gestion/dashboard.html', context)

@login_required
def lista_requerimientos(request):
    if request.method == 'POST':
        titulo = request.POST.get('titulo')
        holding_id = request.POST.get('empresa_holding')
        desglose = request.POST.get('desglose_items')

        if titulo and holding_id and desglose:
            holding_obj = EmpresaHolding.objects.get(id=holding_id)
            nuevo_req = RequerimientoTicket.objects.create(
                solicitante=request.user,
                empresa_holding=holding_obj,
                titulo=titulo,
                desglose_items=desglose
            )
            nuevo_req.ticket_glpi = f'REQ-{nuevo_req.id:04d}'
            nuevo_req.save()

            messages.success(request, '¡Requerimiento registrado exitosamente!')
            return redirect('requerimientos')

    if request.user.is_staff:
        tickets = RequerimientoTicket.objects.select_related('solicitante', 'empresa_holding').all().order_by('-id')
    else:
        tickets = RequerimientoTicket.objects.filter(solicitante=request.user).select_related('empresa_holding').order_by('-id')

    holdings = EmpresaHolding.objects.all()
    return render(request, 'gestion/requerimientos.html', {'tickets': tickets, 'holdings': holdings})

@login_required
def cambiar_estatus_ticket(request, ticket_id, nuevo_estatus):
    if request.user.is_staff:
        ticket = get_object_or_404(RequerimientoTicket, id=ticket_id)
        ticket.estatus = nuevo_estatus
        ticket.save()
        messages.info(request, f'Estatus del Requerimiento #{ticket.ticket_glpi or ticket.id} actualizado.')
    return redirect('requerimientos')

@login_required

def lista_proveedores(request):
    proveedores = Proveedor.objects.all()
    return render(request, 'gestion/proveedores.html', {'proveedores': proveedores})

def exportar_proveedores_pdf(request):
    try:
        from weasyprint import HTML
        proveedores = Proveedor.objects.all().order_by('nombre_empresa')
        html_string = render_to_string('gestion/pdf_proveedores.html', {'proveedores': proveedores})
        pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
        
        response = HttpResponse(pdf_file, content_type='application/pdf')
        response['Content-Disposition'] = 'inline; filename="Lista_de_Proveedores.pdf"'
        return response
    except Exception as e:
        return HttpResponse(f'Error generando PDF: {str(e)}', status=500)

@login_required
def lista_compras(request):
    compras = RegistroCompraPago.objects.prefetch_related(
        'pagos_proveedores__proveedor',
        'proveedor_cot_1', 'proveedor_cot_2', 'proveedor_cot_3'
    ).select_related('empresa_destino').all().order_by('-id')
    
    req_en_curso = RequerimientoTicket.objects.filter(estatus__in=['PENDIENTE', 'EN_CURSO']).order_by('-id')
    return render(request, 'gestion/compras.html', {'compras': compras, 'req_en_curso': req_en_curso})

@login_required
def trazabilidad_facturas(request):
    pagos = PagoProveedorAdjudicado.objects.select_related('solicitud__empresa_destino', 'proveedor').all().order_by('-fecha_pago')
    empresas = EmpresaHolding.objects.all()
    proveedores = Proveedor.objects.all()
    return render(request, 'gestion/facturas.html', {
        'pagos': pagos,
        'empresas': empresas,
        'proveedores': proveedores
    })

@login_required
def lista_inventario(request):
    items = Inventario.objects.select_related('empresa_asignada', 'proveedor').all().order_by('-id')
    empresas = EmpresaHolding.objects.all()
    return render(request, 'gestion/inventario.html', {'items': items, 'empresas': empresas})

from django.http import JsonResponse
from django.views.decorators.csrf import csrf_exempt
from django.shortcuts import get_object_or_404

@csrf_exempt
def transferir_equipo(request, item_id):
    if request.method == 'POST':
        try:
            from gestion.models import Inventario, EmpresaHolding
            item = get_object_or_404(Inventario, id=item_id)
            empresa_id = request.POST.get('empresa_destino')
            if empresa_id:
                empresa_dest = get_object_or_404(EmpresaHolding, id=empresa_id)
                
                # Resguardar empresa_origen si está vacía
                if not item.empresa_origen:
                    item.empresa_origen = item.empresa_asignada or empresa_dest
                
                # Asignar nueva empresa
                item.empresa_asignada = empresa_dest
                item.estatus = 'ASIGNADO'
                item.save()
                
                return JsonResponse({'status': 'ok', 'empresa_nombre': empresa_dest.nombre})
            return JsonResponse({'status': 'error', 'message': 'Por favor selecciona una empresa válida.'}, status=400)
        except Exception as e:
            return JsonResponse({'status': 'error', 'message': str(e)}, status=500)
    return JsonResponse({'status': 'error', 'message': 'Método no permitido'}, status=405)


@login_required
def exportar_inventario_excel(request):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventario_Holding"

    headers = ["ID", "Producto / Modelo", "Serial / Código", "Cantidad", "Proveedor", "Costo Unit. Con IVA ($)", "Costo Total ($)", "Fecha Compra", "Empresa Asignada", "Estatus"]
    
    header_fill = PatternFill(start_color="0D6EFD", end_color="0D6EFD", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center_align = Alignment(horizontal="center", vertical="center")

    ws.append(headers)
    for col_num, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    items = Inventario.objects.select_related('empresa_asignada', 'proveedor').all().order_by('-id')
    for item in items:
        ws.append([
            item.id,
            item.producto_marca_modelo,
            item.serial_limpio,
            item.cantidad,
            item.proveedor.nombre_empresa if item.proveedor else "N/A",
            float(item.costo_unitario_con_iva),
            float(item.costo_total),
            item.fecha_compra.strftime('%d/%m/%Y') if item.fecha_compra else "-",
            item.empresa_asignada.nombre if item.empresa_asignada else "Sin Asignar",
            item.get_estatus_display()
        ])

    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = 'attachment; filename="Reporte_Inventario_Holding.xlsx"'
    wb.save(response)
    return response

@login_required
def exportar_inventario_pdf(request):
    try:
        from weasyprint import HTML
        items = Inventario.objects.select_related('empresa_asignada', 'proveedor').all().order_by('-id')
        try:
            html_string = render_to_string('gestion/pdf_inventario.html', {'items': items})
        except:
            html_string = render_to_string('pdf_inventario.html', {'items': items})
            
        pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
        
        response = HttpResponse(pdf_file, content_type='application/pdf')
        response['Content-Disposition'] = 'inline; filename="Reporte_Inventario_Holding.pdf"'
        return response
    except Exception as e:
        return HttpResponse(f'Error generando PDF: {str(e)}', status=500)

@login_required
def exportar_facturas_pdf(request):
    try:
        from weasyprint import HTML
        pagos = PagoProveedorAdjudicado.objects.select_related('solicitud__empresa_destino', 'proveedor').all().order_by('-fecha_pago')
        try:
            html_string = render_to_string('gestion/pdf_facturas.html', {'pagos': pagos})
        except:
            html_string = render_to_string('pdf_facturas.html', {'pagos': pagos})
            
        pdf_file = HTML(string=html_string, base_url=request.build_absolute_uri('/')).write_pdf()
        
        response = HttpResponse(pdf_file, content_type='application/pdf')
        response['Content-Disposition'] = 'inline; filename="Reporte_Trazabilidad_Facturas.pdf"'
        return response
    except Exception as e:
        return HttpResponse(f'Error generando PDF: {str(e)}', status=500)


from django.contrib.auth.models import User
from django.contrib.auth.decorators import login_required
from django.contrib import messages

@login_required
def usuarios(request):
    if not (request.user.is_superuser or (hasattr(request.user, "perfil") and request.user.perfil.rol == "ADMIN")):
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    if request.method == 'POST':
        username = request.POST.get('username')
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        password = request.POST.get('password')
        rol = request.POST.get('rol', 'SOLICITANTE')

        if User.objects.filter(username=username).exists():
            messages.error(request, f'El usuario "{username}" ya existe.')
        else:
            is_staff = True if rol in ['COMPRAS', 'ADMIN'] else False
            user = User.objects.create_user(username=username, password=password, first_name=first_name, last_name=last_name, is_staff=is_staff)
            perfil, _ = PerfilUsuario.objects.get_or_create(usuario=user)
            perfil.rol = rol
            perfil.save()
            messages.success(request, f'Usuario "{username}" creado exitosamente con el rol {rol}.')
            return redirect('usuarios')

    usuarios_list = User.objects.select_related('perfil').all().order_by('-date_joined')
    return render(request, 'gestion/usuarios.html', {'usuarios_list': usuarios_list})

from django.contrib.auth.models import User
from gestion.models import PerfilUsuario

@login_required
def usuarios(request):
    if not (request.user.is_superuser or (hasattr(request.user, "perfil") and request.user.perfil.rol == "ADMIN")):
        from django.core.exceptions import PermissionDenied
        raise PermissionDenied
    if request.method == 'POST':
        username = request.POST.get('username')
        first_name = request.POST.get('first_name', '')
        last_name = request.POST.get('last_name', '')
        password = request.POST.get('password')
        rol = request.POST.get('rol', 'SOLICITANTE')

        if User.objects.filter(username=username).exists():
            messages.error(request, f'El usuario "{username}" ya existe.')
        else:
            is_staff = True if rol in ['COMPRAS', 'ADMIN'] else False
            user = User.objects.create_user(username=username, password=password, first_name=first_name, last_name=last_name, is_staff=is_staff)
            perfil, _ = PerfilUsuario.objects.get_or_create(usuario=user)
            perfil.rol = rol
            perfil.save()
            messages.success(request, f'Usuario "{username}" creado exitosamente con el rol {rol}.')
            return redirect('usuarios')

    usuarios_list = User.objects.select_related('perfil').all().order_by('-date_joined')
    return render(request, 'gestion/usuarios.html', {'usuarios_list': usuarios_list})

@login_required
def api_obtener_requerimiento(request, ticket_id):
    from django.http import JsonResponse
    from gestion.models import RequerimientoTicket
    try:
        req = RequerimientoTicket.objects.get(pk=ticket_id)
        
        solic_email = ""
        if req.solicitante:
            solic_email = getattr(req.solicitante, 'email', '') or getattr(req.solicitante, 'username', '')

        empresa_id = req.empresa_holding.id if hasattr(req, 'empresa_holding') and req.empresa_holding else ""
        num_ticket = getattr(req, 'ticket_glpi', f'REQ-{req.id:04d}')

        # EXTRAER EL CAMPO EXACTO 'desglose_items'
        items_solicitados = getattr(req, 'desglose_items', '') or ''

        data = {
            'id': req.id,
            'numero_ticket': num_ticket,
            'titulo': getattr(req, 'titulo', ''),
            'persona_solicito': solic_email,
            'empresa_id': empresa_id,
            'fecha': req.creado_el.strftime('%Y-%m-%d') if hasattr(req, 'creado_el') and req.creado_el else '',
            'desglose': items_solicitados
        }
        return JsonResponse(data)
    except Exception as e:
        return JsonResponse({'error': str(e)}, status=400)